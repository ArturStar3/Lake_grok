from io import BytesIO
from zipfile import ZipFile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from docx import Document
from docx.shared import Pt, RGBColor
from rest_framework.exceptions import ValidationError
from rest_framework.test import APITestCase

from accounts.enums import ModuleLevel
from accounts.models import SecurityGroup
from accounts.tests.base import create_admin_group, create_user
from api.demo_scanner import import_docx, normalize_scanner


def document_upload():
    doc = Document()
    doc.add_heading('Источники информации', level=1)
    paragraph = doc.add_paragraph()
    paragraph.alignment = 1
    run = paragraph.add_run('Текст <script> & данные')
    run.bold = True
    run.italic = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Первый'
    table.cell(0, 1).text = 'Второй'
    output = BytesIO()
    doc.save(output)
    return SimpleUploadedFile('report.docx', output.getvalue())


class DemoScannerImportTests(SimpleTestCase):
    def test_formatting_and_table_order(self):
        result = import_docx(document_upload())
        self.assertEqual(result['name'], 'report.docx')
        self.assertEqual([b['type'] for b in result['blocks']], ['paragraph', 'paragraph', 'table'])
        paragraph = result['blocks'][1]
        self.assertEqual(paragraph['style']['textAlign'], 'center')
        run = paragraph['runs'][0]
        self.assertEqual(run['text'], 'Текст <script> & данные')
        self.assertEqual(run['style']['fontWeight'], 'bold')
        self.assertEqual(run['style']['fontStyle'], 'italic')
        self.assertEqual(run['style']['fontSize'], '16pt')
        self.assertEqual(run['style']['color'], '#123456')
        self.assertEqual(result['blocks'][2]['rows'][0][1][0]['runs'][0]['text'], 'Второй')

    def test_invalid_and_empty_files(self):
        for upload in (None, SimpleUploadedFile('bad.docx', b'not a zip'), SimpleUploadedFile('bad.txt', b'text')):
            with self.subTest(upload=upload), self.assertRaises(ValidationError):
                import_docx(upload)
        output = BytesIO()
        Document().save(output)
        with self.assertRaises(ValidationError):
            import_docx(SimpleUploadedFile('empty.docx', output.getvalue()))

    def test_invalid_xml_and_size_limit(self):
        output = BytesIO()
        with ZipFile(output, 'w') as archive:
            archive.writestr('word/document.xml', '<broken')
        with self.assertRaises(ValidationError):
            import_docx(SimpleUploadedFile('broken.docx', output.getvalue()))
        upload = document_upload()
        upload.size = 11 * 1024 * 1024
        with self.assertRaises(ValidationError):
            import_docx(upload)

    def test_normalization_rejects_external_assets_and_active_styles(self):
        result = normalize_scanner({
            'images': [{'src': 'https://external.test/a.png'}, {'src': '/media/../secret'}, {'src': '/media/demo_tableau/a.png'}],
            'scan_min_ms': 3000, 'scan_max_ms': 500, 'effect': 'unknown',
            'document': {'blocks': [{'runs': [{'text': '<img onerror=alert(1)>', 'style': {
                'backgroundImage': 'url(https://external.test)', 'position': 'fixed', 'fontWeight': 'bold',
            }}]}]},
        })
        self.assertEqual(len(result['images']), 1)
        self.assertEqual(result['scan_max_ms'], 3000)
        self.assertEqual(result['effect'], 'letters')
        self.assertEqual(result['document']['blocks'][0]['runs'][0]['style'], {'fontWeight': 'bold'})

    def test_normalization_keeps_configurable_scanner_labels(self):
        result = normalize_scanner({
            'header_left': 'АРХИВ / ОБЗОР',
            'header_right': 'СТАТУС: {status}',
            'paper_meta': 'ДОСЬЕ / {document}',
        })
        self.assertEqual(result['header_left'], 'АРХИВ / ОБЗОР')
        self.assertEqual(result['header_right'], 'СТАТУС: {status}')
        self.assertEqual(result['paper_meta'], 'ДОСЬЕ / {document}')

    def test_malformed_collections_are_safe(self):
        for raw in (None, [], {'images': 42, 'document': {'blocks': 17}}, {'document': {'blocks': [{'runs': 1}, {'type': 'table', 'rows': [1, [1]]}]}}):
            normalize_scanner(raw)


class DemoScannerApiTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        cls.writer = create_user('scanner_writer', groups=[create_admin_group(name='Scanner admins')])
        group = SecurityGroup.objects.create(name='Scanner readers', demo_scenarios=ModuleLevel.READ)
        cls.reader = create_user('scanner_reader', groups=[group])

    def test_import_requires_write_permission(self):
        url = '/api/v1/demo-scenarios/import-scanner-document/'
        self.assertIn(self.client.post(url, {'file': document_upload()}, format='multipart').status_code, (401, 403))
        self.client.force_authenticate(self.reader)
        self.assertEqual(self.client.post(url, {'file': document_upload()}, format='multipart').status_code, 403)
        self.client.force_authenticate(self.writer)
        response = self.client.post(url, {'file': document_upload()}, format='multipart')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['name'], 'report.docx')
        self.assertEqual(self.client.post(url, {}, format='multipart').status_code, 400)

    def test_scanner_scenario_round_trip(self):
        self.client.force_authenticate(self.writer)
        scanner = {'images': [{'src': '/media/demo_tableau/a.png', 'title': 'Источник'}],
                   'document': import_docx(document_upload()), 'effect': 'scramble'}
        response = self.client.post('/api/v1/demo-scenarios/', {
            'title': 'Сканер', 'tableau': {'presets': [{'id': 'scan', 'variant': 'scanner', 'scanner': scanner}]},
            'sequence': [{'type': 'tableau', 'preset_id': 'scan'}],
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        url = f"/api/v1/demo-scenarios/{response.data['id']}/"
        saved = self.client.get(url).data['tableau']['presets'][0]
        self.assertEqual(saved['variant'], 'scanner')
        self.assertEqual(saved['scanner'], normalize_scanner(scanner))
        saved['scanner']['effect'] = 'typewriter'
        response = self.client.patch(url, {'tableau': {'presets': [saved]}}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['tableau']['presets'][0]['scanner']['effect'], 'typewriter')

    def test_network_tableau_keeps_three_logo_links(self):
        self.client.force_authenticate(self.writer)
        network = {
            'logos': [
                {'src': '/media/demo_tableau/logo-1.png', 'title': 'Первый'},
                {'src': '/media/demo_tableau/logo-2.png', 'title': 'Второй'},
                {'src': '/media/demo_tableau/logo-3.png', 'title': 'Третий'},
            ],
        }
        response = self.client.post('/api/v1/demo-scenarios/', {
            'title': 'Сеть',
            'tableau': {'presets': [{'id': 'network', 'variant': 'network', 'network': network}]},
            'sequence': [{'type': 'tableau', 'preset_id': 'network'}],
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        saved = self.client.get(f"/api/v1/demo-scenarios/{response.data['id']}/").data['tableau']['presets'][0]
        self.assertEqual(saved['variant'], 'network')
        self.assertEqual(saved['network'], network)
        saved['network']['logos'][1] = {
            'src': '/media/demo_tableau/replaced-logo.png',
            'title': 'Заменённый',
        }
        response = self.client.patch(
            f"/api/v1/demo-scenarios/{response.data['id']}/",
            {'tableau': {'presets': [saved]}},
            format='json',
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(
            response.data['tableau']['presets'][0]['network']['logos'][1],
            saved['network']['logos'][1],
        )
