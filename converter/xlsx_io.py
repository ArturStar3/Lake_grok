"""Batch xlsx conversion for InfoLake WGS 84 coordinates."""

from __future__ import annotations

import re
from copy import copy
from pathlib import Path
from typing import Iterable, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from coords import (
    SAMPLE_LAT_DD,
    SAMPLE_LAT_DMS,
    SAMPLE_LAT_DMS_RU,
    SAMPLE_LNG_DD,
    SAMPLE_LNG_DMS,
    SAMPLE_LNG_DMS_RU,
    CoordError,
    convert_value,
    format_dd,
    format_dms,
    parse_dd,
    parse_dms,
)

LAT_ALIASES = {
    "lat",
    "latitude",
    "широта",
    "шир",
    "latdd",
    "y",
}
LNG_ALIASES = {
    "lng",
    "lon",
    "long",
    "longitude",
    "долгота",
    "долг",
    "lngdd",
    "x",
}
LAT_DMS_ALIASES = {
    "latdms",
    "latitudedms",
    "широтаdms",
    "dmslat",
}
LNG_DMS_ALIASES = {
    "lngdms",
    "londms",
    "longitudedms",
    "долготаdms",
    "dmslng",
    "dmslon",
}

TEMPLATE_HEADERS = ("lat", "lng", "lat_dms", "lng_dms")


def _norm_header(value: object) -> str:
    text = str(value or "").strip().lower().replace("ё", "е")
    return re.sub(r"[^a-z0-9а-я]+", "", text)


def _classify_header(header: str) -> Optional[str]:
    key = _norm_header(header)
    if not key:
        return None
    if key in LAT_DMS_ALIASES or ("dms" in key and ("lat" in key or "широта" in key)):
        return "lat_dms"
    if key in LNG_DMS_ALIASES or ("dms" in key and ("lng" in key or "lon" in key or "долгота" in key)):
        return "lng_dms"
    if key == "error" or key == "ошибка":
        return "error"
    if key in LAT_ALIASES or key == "широта":
        return "lat"
    if key in LNG_ALIASES or key == "долгота":
        return "lng"
    return None


def detect_columns(headers: list[object]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, header in enumerate(headers):
        role = _classify_header(header)
        if role and role not in mapping:
            mapping[role] = index
    return mapping


def _cell_empty(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def _copy_cell_style(source, target) -> None:
    if source.has_style:
        target.font = copy(source.font)
        target.border = copy(source.border)
        target.fill = copy(source.fill)
        target.number_format = source.number_format
        target.protection = copy(source.protection)
        target.alignment = copy(source.alignment)


def _ensure_column(headers: list[object], mapping: dict[str, int], role: str, title: str) -> int:
    if role in mapping:
        return mapping[role]
    index = len(headers)
    headers.append(title)
    mapping[role] = index
    return index


def _write_row(ws: Worksheet, row_index: int, values: list[object], styles: list) -> None:
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row_index, column=col, value=value)
        if col - 1 < len(styles) and styles[col - 1] is not None:
            _copy_cell_style(styles[col - 1], cell)


def convert_sheet(ws: Worksheet, to_dms: bool) -> tuple[int, int]:
    rows = list(ws.iter_rows(values_only=False))
    if not rows:
        raise CoordError("Пустой лист")

    header_cells = rows[0]
    headers = [cell.value for cell in header_cells]
    styles_by_col = [cell for cell in header_cells]
    mapping = detect_columns(headers)
    if to_dms:
        if "lat" not in mapping or "lng" not in mapping:
            raise CoordError("Не найдены колонки lat/lng (широта/долгота)")
    else:
        has_dms = "lat_dms" in mapping and "lng_dms" in mapping
        has_dd = "lat" in mapping and "lng" in mapping
        if not has_dms and not has_dd:
            raise CoordError("Не найдены колонки DMS или lat/lng")

    lat_dd_i = _ensure_column(headers, mapping, "lat", "lat")
    lng_dd_i = _ensure_column(headers, mapping, "lng", "lng")
    lat_dms_i = _ensure_column(headers, mapping, "lat_dms", "lat_dms")
    lng_dms_i = _ensure_column(headers, mapping, "lng_dms", "lng_dms")
    error_i = _ensure_column(headers, mapping, "error", "error")

    converted = 0
    errors = 0
    out_rows: list[list[object]] = []

    for row_cells in rows[1:]:
        values = [cell.value for cell in row_cells]
        while len(values) < len(headers):
            values.append(None)

        lat_src = values[mapping["lat"]] if "lat" in mapping else None
        lng_src = values[mapping["lng"]] if "lng" in mapping else None
        lat_dms_src = values[mapping["lat_dms"]] if "lat_dms" in mapping else None
        lng_dms_src = values[mapping["lng_dms"]] if "lng_dms" in mapping else None

        if to_dms:
            source_lat, source_lng = lat_src, lng_src
        else:
            source_lat = lat_dms_src if not _cell_empty(lat_dms_src) else lat_src
            source_lng = lng_dms_src if not _cell_empty(lng_dms_src) else lng_src

        if _cell_empty(source_lat) and _cell_empty(source_lng) and all(_cell_empty(v) for v in values[: len(header_cells)]):
            continue

        try:
            if _cell_empty(source_lat) or _cell_empty(source_lng):
                raise CoordError("Нет пары широта/долгота")
            if to_dms:
                lat_dd = parse_dd(source_lat, "lat")
                lng_dd = parse_dd(source_lng, "lng")
            else:
                lat_dd = parse_dms(source_lat, "lat")
                lng_dd = parse_dms(source_lng, "lng")
            values[lat_dd_i] = format_dd(lat_dd) if not to_dms else (values[lat_dd_i] if not _cell_empty(values[lat_dd_i]) else format_dd(lat_dd))
            values[lng_dd_i] = format_dd(lng_dd) if not to_dms else (values[lng_dd_i] if not _cell_empty(values[lng_dd_i]) else format_dd(lng_dd))
            if to_dms:
                values[lat_dms_i] = format_dms(lat_dd, "lat")
                values[lng_dms_i] = format_dms(lng_dd, "lng")
            else:
                if _cell_empty(values[lat_dms_i]):
                    values[lat_dms_i] = format_dms(lat_dd, "lat")
                if _cell_empty(values[lng_dms_i]):
                    values[lng_dms_i] = format_dms(lng_dd, "lng")
            values[error_i] = None
            converted += 1
        except (CoordError, TypeError, ValueError) as exc:
            values[error_i] = str(exc)
            errors += 1

        out_rows.append(values)

    ws.delete_rows(1, ws.max_row)
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        if col - 1 < len(styles_by_col):
            _copy_cell_style(styles_by_col[col - 1], cell)
            cell.font = Font(bold=True)

    for row_index, values in enumerate(out_rows, start=2):
        _write_row(ws, row_index, values, [])

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 22

    return converted, errors


def convert_workbook(path: Path, to_dms: bool, output_path: Optional[Path] = None) -> dict:
    src = Path(path)
    if not src.exists():
        raise FileNotFoundError(src)
    wb = load_workbook(src)
    total_ok = 0
    total_err = 0
    sheet_errors: list[str] = []
    for ws in wb.worksheets:
        if ws.title.strip().lower() == "образец":
            continue
        try:
            ok, err = convert_sheet(ws, to_dms)
            total_ok += ok
            total_err += err
        except CoordError as exc:
            sheet_errors.append(f"{ws.title}: {exc}")
    dest = output_path or src.with_name(f"{src.stem}_converted.xlsx")
    wb.save(dest)
    return {
        "source": str(src),
        "output": str(dest),
        "converted": total_ok,
        "errors": total_err,
        "sheet_errors": sheet_errors,
    }


def convert_files(paths: Iterable[Path], to_dms: bool, output_dir: Optional[Path] = None) -> list[dict]:
    results = []
    for path in paths:
        dest = None
        if output_dir is not None:
            output_dir.mkdir(parents=True, exist_ok=True)
            dest = Path(output_dir) / f"{Path(path).stem}_converted.xlsx"
        results.append(convert_workbook(Path(path), to_dms, dest))
    return results


def _write_sample_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("образец", 0)
    ws["A1"] = "Как писать градусы, минуты, секунды"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:C1")

    rows = [
        (),
        ("Шаблон", 'Г°ММ\'СС.ссс" полушарие', "Г — градусы, ММ — минуты (0–59), СС — секунды (0–59.999)"),
        (),
        ("Что", "Как писать", "Пример"),
        ("Широта (основной вид)", 'Г°ММ\'СС.ссс" N или S', SAMPLE_LAT_DMS),
        ("Долгота (основной вид)", 'Г°ММ\'СС.ссс" E или W', SAMPLE_LNG_DMS),
        ("Без символов ° ' \"", "Г М С полушарие", SAMPLE_LAT_DMS_RU),
        ("Долгота без символов", "Г М С полушарие", SAMPLE_LNG_DMS_RU),
        ("Полушарие впереди", "NГ°ММ'СС.ссс\"", 'N55°45\'04.478"'),
        (),
        ("Полушарие широты", "N или с.ш. = север; S или ю.ш. = юг", ""),
        ("Полушарие долготы", "E или в.д. = восток; W или з.д. = запад", ""),
        ("В файле xlsx", "Колонки lat_dms и lng_dms — по одной координате в ячейке", ""),
        ("InfoLake", "После конвертации lat / lng — десятичные градусы, 6 знаков", f"{SAMPLE_LAT_DD} / {SAMPLE_LNG_DD}"),
    ]
    for index, row in enumerate(rows, start=2):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=index, column=col, value=value)
            if index == 5:
                cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 42
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def save_template(path: Path, to_dms: bool = True) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "coords"
    headers = list(TEMPLATE_HEADERS)
    example = [
        SAMPLE_LAT_DD if to_dms else format_dd(SAMPLE_LAT_DD),
        SAMPLE_LNG_DD if to_dms else format_dd(SAMPLE_LNG_DD),
        convert_value(SAMPLE_LAT_DD, "lat", True),
        convert_value(SAMPLE_LNG_DD, "lng", True),
    ]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        ws.column_dimensions[get_column_letter(col)].width = 24
        ws.cell(row=2, column=col, value=example[col - 1])
    _write_sample_sheet(wb)
    dest = Path(path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest
